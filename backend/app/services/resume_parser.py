import re
import os
from typing import Dict, Any, List, Tuple
import fitz  # PyMuPDF
import docx

from app.models.resume import ParsingStatus
from app.models.profile import SkillStatus, SkillSource

# Curated reference skill keywords across domains for deterministic, grounded extraction
KNOWN_SKILLS = {
    # Languages
    "python", "javascript", "typescript", "c++", "c#", "java", "go", "golang", "rust", "ruby", "php", "swift", "kotlin", "sql", "html", "css", "r", "scala",
    # Frameworks & Libraries
    "fastapi", "django", "flask", "react", "angular", "vue", "next.js", "node.js", "express", "spring boot", "asp.net", "laravel", "tailwind", "bootstrap",
    # Data & AI / ML
    "machine learning", "deep learning", "pytorch", "tensorflow", "scikit-learn", "pandas", "numpy", "opencv", "nlp", "large language models", "llm", "langchain", "langgraph", "rag", "data science",
    # Databases & Caching
    "postgresql", "mysql", "sqlite", "mongodb", "redis", "elasticsearch", "neo4j", "cassandra", "dynamodb",
    # DevOps, Cloud & Tools
    "docker", "kubernetes", "aws", "gcp", "azure", "git", "github", "gitlab", "ci/cd", "linux", "terraform", "ansible", "nginx", "graphql", "rest api", "microservices"
}

EMAIL_REGEX = re.compile(r"[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+")
PHONE_REGEX = re.compile(r"(\+?\d{1,3}[-.\s]?)?(\(?\d{3}\)?[-.\s]?)?\d{3}[-.\s]?\d{4}")
URL_REGEX = re.compile(r"(https?://[^\s]+|www\.[^\s]+)")


class ResumeParserService:
    @staticmethod
    def extract_text(file_path: str, ext: str) -> Tuple[str, Dict[str, Any]]:
        text_content = ""
        metadata: Dict[str, Any] = {}

        if ext == ".pdf":
            try:
                doc = fitz.open(file_path)
                metadata["page_count"] = len(doc)
                metadata["author"] = doc.metadata.get("author", "")
                metadata["title"] = doc.metadata.get("title", "")
                
                pages_text = []
                for page_num in range(len(doc)):
                    page = doc.load_page(page_num)
                    pages_text.append(page.get_text())
                doc.close()
                text_content = "\n".join(pages_text)
            except Exception as e:
                text_content = ""
                metadata["error"] = f"PyMuPDF error: {str(e)}"

        elif ext == ".docx":
            try:
                doc = docx.Document(file_path)
                paragraphs_text = [p.text for p in doc.paragraphs if p.text.strip()]
                # Also include text from tables if any
                table_text = []
                for table in doc.tables:
                    for row in table.rows:
                        row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                        if row_cells:
                            table_text.append(" | ".join(row_cells))
                
                text_content = "\n".join(paragraphs_text + table_text)
                metadata["paragraph_count"] = len(paragraphs_text)
            except Exception as e:
                text_content = ""
                metadata["error"] = f"python-docx error: {str(e)}"
        
        return text_content.strip(), metadata

    @staticmethod
    def parse_resume(file_path: str, original_filename: str) -> Dict[str, Any]:
        ext = os.path.splitext(file_path)[1].lower()
        raw_text, raw_meta = ResumeParserService.extract_text(file_path, ext)

        result: Dict[str, Any] = {
            "parsed_text": raw_text,
            "raw_metadata": raw_meta,
            "parsing_status": ParsingStatus.PARSED,
            "parsing_confidence": 1.0,
            "parsing_notes": "Deterministic extraction complete with evidence tracking.",
            "extracted_skills": [],
            "extracted_education": [],
            "extracted_experience": [],
            "extracted_projects": [],
            "contact_info": {}
        }

        # Check for empty or unselectable text
        if not raw_text or len(raw_text.strip()) < 80:
            result["parsing_status"] = ParsingStatus.PARTIAL
            result["parsing_confidence"] = 0.25
            result["parsing_notes"] = (
                "Insufficient extractable text detected. This document may be a scanned image or use "
                "custom embedded fonts without Unicode mappings. No facts were invented."
            )
            return result

        # 1. Contact Info & URLs
        emails = EMAIL_REGEX.findall(raw_text)
        phones = PHONE_REGEX.findall(raw_text)
        urls = URL_REGEX.findall(raw_text)
        
        result["contact_info"] = {
            "email": emails[0] if emails else None,
            "phone": phones[0][0] if phones and isinstance(phones[0], tuple) else None,
            "links": list(set(urls))[:5]
        }

        # 2. Extract Skills with Exact Evidence
        lines = raw_text.splitlines()
        found_skills = {}
        lower_text = raw_text.lower()

        for skill_kw in KNOWN_SKILLS:
            # Word-boundary search to avoid false positives (e.g., 'go' inside 'good')
            pattern = r"\b" + re.escape(skill_kw) + r"\b"
            match = re.search(pattern, lower_text)
            if match:
                # Find line with evidence
                evidence_line = ""
                for line in lines:
                    if skill_kw in line.lower():
                        evidence_line = line.strip()[:160]
                        break
                
                # Determine display name
                display_name = skill_kw.title() if len(skill_kw) > 3 else skill_kw.upper()
                if skill_kw in ["fastapi", "next.js", "node.js"]:
                    display_name = skill_kw
                elif skill_kw == "ci/cd":
                    display_name = "CI/CD"

                found_skills[skill_kw] = {
                    "name": display_name,
                    "category": ResumeParserService._categorize_skill(skill_kw),
                    "status": SkillStatus.VERIFIED,
                    "source": SkillSource.RESUME_PARSED,
                    "evidence": f"Found in '{original_filename}': \"{evidence_line}\"" if evidence_line else f"Found in resume text",
                    "confidence_score": 1.0
                }

        result["extracted_skills"] = list(found_skills.values())

        # 3. Section Slicing (Education, Experience, Projects)
        sections = ResumeParserService._split_sections(lines)
        result["extracted_education"] = sections.get("education", [])
        result["extracted_experience"] = sections.get("experience", [])
        result["extracted_projects"] = sections.get("projects", [])

        # Uncertainty calculation based on text extraction quality
        if len(result["extracted_skills"]) == 0:
            result["parsing_confidence"] = 0.65
            result["parsing_notes"] = "Text extracted cleanly, but standard technical skill keywords were sparse or non-standard."
        else:
            result["parsing_confidence"] = 0.95

        return result

    @staticmethod
    def _categorize_skill(skill: str) -> str:
        skill = skill.lower()
        if skill in {"python", "javascript", "typescript", "c++", "c#", "java", "go", "golang", "rust", "ruby", "php", "sql", "html", "css"}:
            return "programming_languages"
        elif skill in {"fastapi", "django", "flask", "react", "angular", "vue", "next.js", "node.js", "express", "spring boot"}:
            return "frameworks"
        elif skill in {"machine learning", "deep learning", "pytorch", "tensorflow", "nlp", "llm", "langchain", "langgraph", "rag"}:
            return "ai_data"
        elif skill in {"docker", "kubernetes", "aws", "gcp", "azure", "git", "ci/cd", "linux"}:
            return "tools_infrastructure"
        elif skill in {"postgresql", "mysql", "sqlite", "mongodb", "redis"}:
            return "databases"
        return "technical"

    @staticmethod
    def _split_sections(lines: List[str]) -> Dict[str, List[Dict[str, Any]]]:
        sections: Dict[str, List[Dict[str, Any]]] = {
            "education": [],
            "experience": [],
            "projects": []
        }

        current_section = None
        current_buffer = []

        section_headers = {
            "education": ["education", "academic background", "academics"],
            "experience": ["experience", "work experience", "employment history", "professional experience"],
            "projects": ["projects", "personal projects", "academic projects", "key projects"]
        }

        for line in lines:
            line_str = line.strip()
            if not line_str:
                continue

            lower_line = line_str.lower()
            matched_header = None
            for sec_name, keywords in section_headers.items():
                if any(lower_line == kw or lower_line.startswith(kw + ":") for kw in keywords):
                    matched_header = sec_name
                    break

            if matched_header:
                current_section = matched_header
                continue

            if current_section and len(line_str) > 3:
                # Capture informative bullets or entries
                if current_section == "education" and any(deg in lower_line for deg in ["bachelor", "master", "b.tech", "m.tech", "bs", "ms", "phd", "university", "college", "institute"]):
                    sections["education"].append({
                        "raw_entry": line_str,
                        "source": "resume_section"
                    })
                elif current_section == "experience" and len(sections["experience"]) < 8:
                    if line_str.startswith(("-", "•", "*")) or any(role in lower_line for role in ["engineer", "developer", "lead", "intern", "manager"]):
                        sections["experience"].append({
                            "raw_entry": line_str,
                            "source": "resume_section"
                        })
                elif current_section == "projects" and len(sections["projects"]) < 8:
                    if line_str.startswith(("-", "•", "*")) or ":" in line_str:
                        sections["projects"].append({
                            "raw_entry": line_str,
                            "source": "resume_section"
                        })

        return sections
